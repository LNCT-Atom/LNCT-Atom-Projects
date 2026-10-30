"""Comprehensive Automated Tests for Resume Parsing System satisfying Tests A through G.

Test A — Complete Text Extraction (PDF multi-page & DOCX paragraphs/tables)
Test B — Deterministic Extraction (Profile, Education, Skills, Certs, Achievements)
Test C — AI Boundary (Mocked OpenAI client, strict isolation of Projects/Experience)
Test D — JSON Merge (Deterministic + AI merged into canonical schema)
Test E — MongoDB Persistence (1 complete document, real/mock ID returned)
Test F — Frontend Contract & Persistence Validation
Test G — End-to-End Regression Test (All 7 sections, multiple entries, student without internship)
"""
import io
import json
import pytest
import fitz  # PyMuPDF
from docx import Document
from unittest.mock import MagicMock, patch

from app.schemas.resume_schema import (
    ResumeCanonicalSchema,
    ResumeMetadataSchema,
    ProfileSchema,
    EducationItemSchema,
    TechnicalSkillsSchema,
    ProjectItemSchema,
    ExperienceItemSchema,
    CertificationItemSchema,
    AchievementItemSchema,
    ExperienceType,
    ParseResumeResponse,
)
from app.services.text_extractor import TextExtractor, TextExtractionError
from app.services.section_detector import SectionDetector
from app.services.deterministic_parser import DeterministicParser
from app.services.ai_extractor import AIExtractor, AIExtractionError
from app.services.resume_service import ResumeService
from app.repositories.resume_repository import ResumeRepository, ResumeRepositoryError


# ==============================================================================
# TEST A: Complete Text Extraction (PDF & DOCX with tables & multi-page)
# ==============================================================================
def test_test_a_pdf_multi_page_text_extraction():
    """Verify PyMuPDF extracts text across multiple pages into the final output."""
    doc = fitz.open()
    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 72), "Rahul Sharma\nEmail: rahul.sharma@example.com\nEDUCATION\nB.Tech CSE")
    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 72), "PROJECTS\nCampus Placement Portal\nWORK EXPERIENCE\nSoftware Engineer")

    pdf_bytes = doc.write()
    doc.close()

    res = TextExtractor.extract_from_pdf(pdf_bytes)
    extracted_text, is_scanned, page_count = res

    assert page_count == 2
    assert not is_scanned
    # Both page 1 and page 2 text must be present
    assert "Rahul Sharma" in extracted_text
    assert "B.Tech CSE" in extracted_text
    assert "Campus Placement Portal" in extracted_text
    assert "Software Engineer" in extracted_text


def test_test_a_docx_paragraphs_and_tables_extraction():
    """Verify python-docx extracts text from both body paragraphs and tables."""
    doc = Document()
    doc.add_paragraph("Ananya Verma")
    doc.add_paragraph("Email: ananya@example.com | Phone: +91 9876543210")
    doc.add_paragraph("EDUCATION")

    # Add a table for education
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Degree"
    hdr_cells[1].text = "Institution"
    hdr_cells[2].text = "Passing Year"
    hdr_cells[3].text = "CGPA"

    row1 = table.add_row().cells
    row1[0].text = "B.Tech CSE"
    row1[1].text = "National Institute of Technology Karnataka"
    row1[2].text = "2020 - 2024"
    row1[3].text = "8.95"

    row2 = table.add_row().cells
    row2[0].text = "Class XII"
    row2[1].text = "Delhi Public School"
    row2[2].text = "2019"
    row2[3].text = "95%"

    doc.add_paragraph("TECHNICAL SKILLS")
    doc.add_paragraph("Languages: Python, TypeScript, Java")

    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    extracted_text, is_scanned, _ = TextExtractor.extract_from_docx(docx_bytes)

    assert "Ananya Verma" in extracted_text
    assert "ananya@example.com" in extracted_text
    # Table data must be preserved with cell separators
    assert "National Institute of Technology Karnataka" in extracted_text
    assert "Delhi Public School" in extracted_text
    assert "8.95" in extracted_text
    assert "Python, TypeScript" in extracted_text


# ==============================================================================
# TEST B: Deterministic Extraction (Profile, Education, Skills, Certs, Achievements)
# ==============================================================================
def test_test_b_deterministic_extraction_all_five_sections():
    """Verify Profile, Education, Skills, Certs, Achievements are independently populated."""
    header = """
    Vikram Malhotra
    vikram.m@example.com | +91 9123456789
    Bengaluru, India
    linkedin.com/in/vikram-malhotra
    github.com/vikram-m
    https://vikram.dev
    """
    summary = "Aspiring Full Stack Engineer passionate about cloud-native distributed systems."

    edu_text = """
    B.Tech in Computer Science and Engineering
    Delhi Technological University
    2020 - 2024 | CGPA: 9.15 / 10

    Class XII (Senior Secondary)
    Delhi Public School, R.K. Puram
    2018 - 2019 | Percentage: 94.8%

    Class X (Secondary School)
    Delhi Public School, R.K. Puram
    2016 - 2017 | Percentage: 96.2%
    """

    skills_text = """
    Programming Languages: Python, Java, TypeScript, Go
    Frameworks & Web: React, FastAPI, Node.js, Spring Boot
    Databases: PostgreSQL, MongoDB, Redis
    Tools & DevOps: Docker, Kubernetes, Git, AWS
    Other Skills: Microservices, System Design, REST APIs
    """

    certs_text = """
    AWS Certified Solutions Architect - Associate
    Issued by Amazon Web Services | Valid: Aug 2023 | Credential ID: AWS-789012
    https://aws.amazon.com/verify/AWS-789012

    Postman API Fundamentals Student Expert
    Issued by Postman - Dec 2022
    """

    achs_text = """
    1st Place Winner - Smart India Hackathon 2023: Developed an automated triage platform for healthcare workers.
    ACM ICPC Regional Finalist 2023: Ranked among top 30 teams nationally.
    """

    full_text = f"{header}\nSUMMARY\n{summary}\nEDUCATION\n{edu_text}\nSKILLS\n{skills_text}\nCERTIFICATIONS\n{certs_text}\nACHIEVEMENTS\n{achs_text}"

    # 1. Profile
    profile = DeterministicParser.parse_profile(header, summary, full_text)
    assert profile.full_name == "Vikram Malhotra"
    assert profile.email == "vikram.m@example.com"
    assert profile.phone == "+91 9123456789"
    assert "Bengaluru" in profile.location
    assert "linkedin.com/in/vikram-malhotra" in profile.linkedin_url
    assert "github.com/vikram-m" in profile.github_url
    assert profile.portfolio_url == "https://vikram.dev"
    assert profile.summary == summary

    # 2. Education (3 distinct qualifications, CGPA & Percentage separated)
    edu_items = DeterministicParser.parse_education(edu_text)
    assert len(edu_items) == 3

    btech = next(e for e in edu_items if e.degree == "B.Tech")
    assert "Delhi Technological University" in btech.institution
    assert "Computer Science" in btech.specialization
    assert btech.cgpa == 9.15
    assert btech.cgpa_scale == 10.0
    assert btech.percentage is None  # Strict separation!
    assert btech.start_year == "2020"
    assert btech.end_year == "2024"

    class12 = next(e for e in edu_items if "Class XII" in (e.degree or ""))
    assert "Delhi Public School" in class12.institution
    assert class12.percentage == "94.8%"
    assert class12.cgpa is None  # Strict separation!

    class10 = next(e for e in edu_items if "Class X /" in (e.degree or ""))
    assert class10.percentage == "96.2%"

    # 3. Technical Skills
    skills = DeterministicParser.parse_technical_skills(skills_text)
    assert "Python" in skills.programming_languages
    assert "TypeScript" in skills.programming_languages
    assert "React" in skills.frameworks
    assert "FastAPI" in skills.frameworks
    assert "PostgreSQL" in skills.databases
    assert "MongoDB" in skills.databases
    assert "Docker" in skills.tools
    assert "Kubernetes" in skills.tools
    assert any("System Design" in s for s in skills.other_skills)

    # 4. Certifications
    certs = DeterministicParser.parse_certifications(certs_text)
    assert len(certs) == 2
    assert "AWS Certified Solutions Architect" in certs[0].name
    assert certs[0].issuer == "Amazon Web Services"
    assert certs[0].credential_id == "AWS-789012"
    assert certs[1].issuer == "Postman"

    # 5. Achievements
    achs = DeterministicParser.parse_achievements(achs_text)
    assert len(achs) == 2
    assert "Winner" in achs[0].title
    assert "Smart India Hackathon" in achs[0].description


# ==============================================================================
# TEST C: AI Boundary (Mocked OpenAI client, strictly Projects & Experience only)
# ==============================================================================
@patch("openai.resources.chat.completions.Completions.create")
def test_test_c_ai_boundary_strictly_projects_and_experience_only(mock_create):
    """Verify that OpenAI ONLY receives isolated Projects & Experience text.

    Assert that the 5 deterministic sections and the full resume are strictly excluded.
    """
    mock_choice = MagicMock()
    mock_choice.message.content = json.dumps({
        "projects": [
            {
                "title": "Distributed Task Queue",
                "role": "Backend Lead",
                "description": "Engineered worker queue with Redis and Python.",
                "technical_skills": ["Python", "Redis"]
            }
        ],
        "experience": [
            {
                "company": "EPAM Systems",
                "role": "Software Engineering Intern",
                "experience_type": "internship",
                "start_date": "Jan 2024",
                "end_date": "Jun 2024",
                "description": "Developed microservices and wrote unit tests.",
                "technical_skills": ["Java", "Docker"]
            }
        ]
    })
    mock_create.return_value = MagicMock(choices=[mock_choice])

    with patch("app.core.config.settings.OPENAI_API_KEY", "mock-test-key"):
        projs, exp, ai_called = AIExtractor.extract_projects_and_experience(
            projects_text="Distributed Task Queue in Python and Redis",
            experience_text="EPAM Systems Intern Jan 2024 - Jun 2024",
        )

    assert ai_called is True
    assert len(projs) == 1
    assert projs[0].title == "Distributed Task Queue"
    assert len(exp) == 1
    assert exp[0].company == "EPAM Systems"
    assert exp[0].experience_type == ExperienceType.INTERNSHIP

    # Check payload sent to OpenAI
    call_args = mock_create.call_args[1]
    messages = call_args["messages"]
    sent_prompt = messages[1]["content"]

    # Must contain ONLY projects and experience
    assert "ISOLATED PROJECTS SECTION TEXT" in sent_prompt
    assert "ISOLATED EXPERIENCE SECTION TEXT" in sent_prompt

    # Must NEVER contain other sections
    assert "EDUCATION" not in sent_prompt
    assert "PROFILE" not in sent_prompt
    assert "SKILLS" not in sent_prompt
    assert "CERTIFICATIONS" not in sent_prompt
    assert "ACHIEVEMENTS" not in sent_prompt


# ==============================================================================
# TEST D: JSON Merge (Deterministic + AI results into canonical schema)
# ==============================================================================
def test_test_d_json_merge_preserves_all_sections():
    """Verify deterministic + AI outputs merge into one canonical schema without dropping or renaming fields."""
    profile = ProfileSchema(
        full_name="Meera Iyer",
        email="meera@example.com",
        phone="+91 9988776655",
        location="Pune, India",
    )
    education = [
        EducationItemSchema(
            institution="College of Engineering Pune",
            degree="B.Tech",
            specialization="Information Technology",
            start_year="2019",
            end_year="2023",
            cgpa="8.75",
        )
    ]
    skills = TechnicalSkillsSchema(
        programming_languages=["Python", "C++"],
        frameworks=["FastAPI"],
        databases=["MongoDB"],
        tools=["Git", "Docker"],
        other_skills=["Agile"],
    )
    projects = [
        ProjectItemSchema(
            title="Smart Attendance System",
            role="Creator",
            description="Facial recognition attendance system.",
            technical_skills=["Python", "OpenCV"],
        )
    ]
    experience = [
        ExperienceItemSchema(
            company="TCS",
            role="Systems Engineer",
            experience_type=ExperienceType.FULL_TIME,
            start_date="Aug 2023",
            end_date="Present",
            description="Built microservices.",
            technical_skills=["Python", "FastAPI"],
        )
    ]
    certifications = [
        CertificationItemSchema(
            name="Azure Fundamentals AZ-900",
            issuer="Microsoft",
            issue_date="2023",
        )
    ]
    achievements = [
        AchievementItemSchema(
            title="Best Capstone Project Award",
            description="Awarded 1st place in university exhibition.",
        )
    ]
    metadata = ResumeMetadataSchema(
        parser_version="1.0",
        status="completed",
        file_name="meera_resume.pdf",
        file_hash="mockhash123",
    )

    canonical = ResumeCanonicalSchema(
        profile=profile,
        education=education,
        technical_skills=skills,
        projects=projects,
        experience=experience,
        certifications=certifications,
        achievements=achievements,
        metadata=metadata,
    )

    data_dict = canonical.model_dump()

    # Verify top-level JSON structure exactly matches the canonical contract
    assert "profile" in data_dict
    assert "education" in data_dict
    assert "technical_skills" in data_dict
    assert "projects" in data_dict
    assert "experience" in data_dict
    assert "certifications" in data_dict
    assert "achievements" in data_dict
    assert "metadata" in data_dict

    assert data_dict["profile"]["full_name"] == "Meera Iyer"
    assert data_dict["education"][0]["degree"] == "B.Tech"
    assert "Python" in data_dict["technical_skills"]["technologies"]
    assert "C++" in data_dict["technical_skills"]["technologies"]
    assert data_dict["projects"][0]["title"] == "Smart Attendance System"
    assert data_dict["experience"][0]["company"] == "TCS"
    assert data_dict["certifications"][0]["issuer"] == "Microsoft"
    assert data_dict["achievements"][0]["title"] == "Best Capstone Project Award"


# ==============================================================================
# TEST E: MongoDB Persistence (Complete document persisted, ID returned)
# ==============================================================================
def test_test_e_mongodb_persistence_and_failure_handling():
    """Verify one complete document containing all 7 sections is persisted and document ID returned."""
    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.find_by_file_hash.return_value = None
    mock_repo.insert_resume.return_value = "664fa1e2b58872f23b2075ad"

    service = ResumeService(repository=mock_repo)

    with patch.object(TextExtractor, "extract_text", return_value={
        "text": "Deepak Kumar\ndeepak@example.com\nEDUCATION\nB.Tech CSE\nNIT Trichy",
        "is_scanned": False,
        "page_count": 1,
        "file_hash": "hash_deepak_123",
        "file_size": 2048,
    }):
        response = service.process_resume("deepak_resume.pdf", b"%PDF-mock")

        # Verify response contract
        assert response.success is True
        assert response.saved is True
        assert response.resume_id == "664fa1e2b58872f23b2075ad"
        assert response.status == "completed"
        assert "submitted successfully" in response.message

        # Verify that insert_resume was called with all 7 sections in 1 document
        call_doc = mock_repo.insert_resume.call_args[0][0]
        assert "profile" in call_doc
        assert "education" in call_doc
        assert "technical_skills" in call_doc
        assert "projects" in call_doc
        assert "experience" in call_doc
        assert "certifications" in call_doc
        assert "achievements" in call_doc
        assert "metadata" in call_doc


def test_test_e_mongodb_persistence_failure_raises_error():
    """Verify that if database persistence fails, an error is raised and no success is reported."""
    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.find_by_file_hash.return_value = None
    mock_repo.insert_resume.side_effect = ResumeRepositoryError("MongoDB connection timeout")

    service = ResumeService(repository=mock_repo)

    with patch.object(TextExtractor, "extract_text", return_value={
        "text": "Deepak Kumar\ndeepak@example.com\nEDUCATION\nB.Tech CSE",
        "is_scanned": False,
        "page_count": 1,
        "file_hash": "hash_deepak_fail",
        "file_size": 2048,
    }):
        with pytest.raises(ResumeRepositoryError, match="Database persistence failed"):
            service.process_resume("deepak_resume.pdf", b"%PDF-mock")


# ==============================================================================
# TEST F: Frontend Contract Validation
# ==============================================================================
def test_test_f_backend_response_contract_guarantees_persistence():
    """Verify that ParseResumeResponse strictly enforces persistence confirmation fields."""
    canonical = ResumeCanonicalSchema(
        profile=ProfileSchema(full_name="Pooja Sharma"),
        education=[EducationItemSchema(degree="B.Tech", institution="IIT Delhi")],
    )

    resp = ParseResumeResponse(
        success=True,
        message="Your resume has been submitted successfully, and your parsed JSON data has been saved to our database.",
        saved=True,
        resume_id="507f1f77bcf86cd799439011",
        data=canonical,
    )

    resp_dict = resp.model_dump()
    assert resp_dict["success"] is True
    assert resp_dict["saved"] is True
    assert resp_dict["resume_id"] == "507f1f77bcf86cd799439011"
    assert resp_dict["document_id"] == "507f1f77bcf86cd799439011"
    assert "saved to our database" in resp_dict["message"]


# ==============================================================================
# TEST G: Regression Test (Representative resume with all 7 sections, multiple entries)
# ==============================================================================
def test_test_g_end_to_end_student_without_internship_all_sections():
    """Regression Test: Student with no internship experience, multiple education entries,

    multiple projects, certifications, and achievements.
    Verify all 7 sections are preserved and stored in the canonical schema.
    """
    resume_text = """
    Aarav Patel
    aarav.patel@example.com | +91 9876543210
    Mumbai, Maharashtra, India
    linkedin.com/in/aarav.patel
    github.com/aaravpatel
    https://aaravpatel.dev

    PROFESSIONAL SUMMARY
    Pre-final year Computer Engineering student with a strong background in backend systems and distributed data structures.

    EDUCATION
    B.Tech in Computer Engineering
    Veermata Jijabai Technological Institute, Mumbai
    2021 - 2025 | CGPA: 9.35 / 10

    Class XII (Higher Secondary Certificate)
    R.D. National College, Mumbai
    2019 - 2021 | Percentage: 92.4%

    Class X (Secondary School Certificate)
    St. Mary's High School, Mumbai
    2018 - 2019 | Percentage: 95.0%

    TECHNICAL SKILLS
    Languages: Python, C++, JavaScript, Go, SQL
    Frameworks: FastAPI, React, Node.js, Express
    Databases: MongoDB, PostgreSQL, Redis
    Tools: Git, Docker, Kubernetes, Linux, Postman
    Other Skills: Microservices, System Architecture, Algorithms

    PROJECTS
    Real-Time Collaborative Code Editor
    Built with React, Node.js, and WebSocket. Enabled 50+ concurrent users to edit code in real time.

    Automated Resume Parsing Platform
    Engineered a dual-pipeline parser in FastAPI extracting deterministic sections and AI-enhanced projects.

    CERTIFICATIONS
    AWS Certified Cloud Practitioner
    Issued by Amazon Web Services - Oct 2023 | Credential ID: CLD-102938

    MongoDB Certified Developer Associate
    Issued by MongoDB - Jan 2024

    ACHIEVEMENTS
    Winner - Mumbai University Hackathon 2023: First place among 120 participating engineering teams.
    CodeChef 4-Star Coder: Max rating 1845.
    """

    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.find_by_file_hash.return_value = None
    mock_repo.insert_resume.return_value = "doc_aarav_final_664fa"

    service = ResumeService(repository=mock_repo)

    # Mock OpenAI chat completion for the 2 projects (student has NO internship)
    with patch("openai.resources.chat.completions.Completions.create") as mock_openai, \
         patch("app.core.config.settings.OPENAI_API_KEY", "mock-ai-key"):

        mock_choice = MagicMock()
        mock_choice.message.content = json.dumps({
            "projects": [
                {
                    "title": "Real-Time Collaborative Code Editor",
                    "role": "Creator",
                    "description": "Built with React, Node.js, and WebSocket for 50+ concurrent users.",
                    "technical_skills": ["React", "Node.js", "WebSocket"]
                },
                {
                    "title": "Automated Resume Parsing Platform",
                    "role": "Backend Developer",
                    "description": "Dual-pipeline parser in FastAPI.",
                    "technical_skills": ["FastAPI", "Python", "MongoDB"]
                }
            ],
            "experience": []  # Student has NO internship experience!
        })
        mock_openai.return_value = MagicMock(choices=[mock_choice])

        with patch.object(TextExtractor, "extract_text", return_value={
            "text": resume_text,
            "is_scanned": False,
            "page_count": 1,
            "file_hash": "hash_aarav_regression_test",
            "file_size": 3500,
        }):
            response = service.process_resume("aarav_resume.pdf", b"%PDF-mock")

    # 1. Persistence & Contract Assertions
    assert response.success is True
    assert response.saved is True
    assert response.resume_id == "doc_aarav_final_664fa"
    assert "saved to our database" in response.message

    data = response.data

    # 2. Profile assertions
    assert data.profile.full_name == "Aarav Patel"
    assert data.profile.email == "aarav.patel@example.com"
    assert data.profile.phone == "+91 9876543210"
    assert "Mumbai" in data.profile.location
    assert "linkedin.com/in/aarav.patel" in data.profile.linkedin_url
    assert "github.com/aaravpatel" in data.profile.github_url
    assert data.profile.portfolio_url == "https://aaravpatel.dev"
    assert "Pre-final year" in (data.profile.summary or "")

    # 3. Education assertions (Multiple entries, strict CGPA & Percentage separation)
    assert len(data.education) == 3
    btech = next(e for e in data.education if e.degree == "B.Tech")
    assert "Veermata Jijabai Technological Institute" in btech.institution
    assert btech.cgpa == 9.35
    assert btech.cgpa_scale == 10.0
    assert btech.percentage is None
    assert btech.start_year == "2021"
    assert btech.end_year == "2025"

    class12 = next(e for e in data.education if "Class XII" in (e.degree or ""))
    assert "R.D. National College" in class12.institution
    assert class12.percentage == "92.4%"
    assert class12.cgpa is None

    class10 = next(e for e in data.education if "Class X /" in (e.degree or ""))
    assert "St. Mary's High School" in class10.institution
    assert class10.percentage == "95.0%"
    assert class10.cgpa is None

    # 4. Technical skills assertions
    assert "Python" in data.technical_skills.programming_languages
    assert "Go" in data.technical_skills.programming_languages
    assert "FastAPI" in data.technical_skills.frameworks
    assert "MongoDB" in data.technical_skills.databases
    assert "Docker" in data.technical_skills.tools

    # 5. Projects assertions (Extracted via AI)
    assert len(data.projects) == 2
    assert data.projects[0].title == "Real-Time Collaborative Code Editor"
    assert "WebSocket" in data.projects[0].technical_skills

    # 6. Experience assertions (Student without internship -> empty array, NOT null, NOT hallucinated)
    assert data.experience == []

    # 7. Certifications assertions
    assert len(data.certifications) == 2
    assert "AWS Certified Cloud Practitioner" in data.certifications[0].name
    assert data.certifications[0].issuer == "Amazon Web Services"

    # 8. Achievements assertions
    assert len(data.achievements) == 2
    assert "Winner" in data.achievements[0].title

    assert data.metadata.file_name == "aarav_resume.pdf"
    assert data.metadata.ai_called_for_projects_or_experience is True
    assert data.metadata.status == "completed"


# ==============================================================================
# ADDITIONAL VALIDATION & SECURITY TESTS
# ==============================================================================
def test_validate_empty_file_rejected():
    with pytest.raises(TextExtractionError, match="empty"):
        TextExtractor.validate_file("empty.pdf", b"")


def test_validate_unsupported_file_extension():
    with pytest.raises(TextExtractionError, match="Unsupported file type"):
        TextExtractor.validate_file("resume.exe", b"test content")


def test_validate_corrupted_signature_pdf():
    with pytest.raises(TextExtractionError, match="Missing standard %PDF signature"):
        TextExtractor.validate_file("resume.pdf", b"This is not a pdf file")


def test_validate_file_size_limit():
    large_bytes = b"%PDF" + b"x" * (11 * 1024 * 1024)
    with pytest.raises(TextExtractionError, match="exceeds maximum allowed limit"):
        TextExtractor.validate_file("large.pdf", large_bytes, max_size_mb=10)


def test_duplicate_resume_detection_when_demo_mode_disabled():
    """Verify that when DEMO_MODE is False, duplicate detection returns existing record."""
    mock_repo = MagicMock(spec=ResumeRepository)
    fake_existing_doc = {
        "_id": "507f1f77bcf86cd799439011",
        "profile": {"full_name": "Existing Candidate", "email": "test@example.com"},
        "education": [],
        "technical_skills": {
            "programming_languages": [],
            "frameworks": [],
            "databases": [],
            "tools": [],
            "other_skills": []
        },
        "projects": [],
        "experience": [],
        "certifications": [],
        "achievements": [],
        "metadata": {"parser_version": "1.0", "status": "completed", "file_hash": "abc"}
    }
    mock_repo.find_by_file_hash.return_value = fake_existing_doc

    service = ResumeService(repository=mock_repo)

    with patch("app.core.config.settings.DEMO_MODE", False), \
         patch.object(TextExtractor, "extract_text", return_value={
             "text": "Existing Candidate\ntest@example.com",
             "is_scanned": False,
             "page_count": 1,
             "file_hash": "abc",
             "file_size": 100
         }):
        resp = service.process_resume("resume.pdf", b"%PDF-dummy")
        assert resp.duplicate is True
        assert resp.resume_id == "507f1f77bcf86cd799439011"
        assert resp.data.profile.full_name == "Existing Candidate"


def test_demo_mode_allows_repeated_uploads():
    """Verify that in DEMO_MODE (default True), duplicate detection is bypassed and every upload creates a new record."""
    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.insert_resume.return_value = "new_demo_doc_id_999"

    service = ResumeService(repository=mock_repo)

    with patch("app.core.config.settings.DEMO_MODE", True), \
         patch.object(TextExtractor, "extract_text", return_value={
             "text": "Repeat Candidate\nrepeat@example.com\nEDUCATION\nB.Tech CSE",
             "is_scanned": False,
             "page_count": 1,
             "file_hash": "same_hash_uploaded_twice",
             "file_size": 1500
         }):
        resp = service.process_resume("demo_resume.pdf", b"%PDF-mock")
        # In demo mode, duplicate warning is false and document is re-parsed and inserted
        assert resp.duplicate is False
        assert resp.resume_id == "new_demo_doc_id_999"
        assert resp.success is True
        assert resp.saved is True
        assert mock_repo.insert_resume.called


def test_repository_list_and_delete_resumes():
    """Verify listing and deleting resumes in repository."""
    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.list_resumes.return_value = [
        {
            "id": "doc123",
            "filename": "candidate.pdf",
            "created_at": "2026-10-23T00:30:00Z",
            "status": "completed",
            "candidate_name": "Test User",
            "candidate_email": "test@example.com",
        }
    ]
    mock_repo.delete_resume.return_value = True

    items = mock_repo.list_resumes(limit=10)
    assert len(items) == 1
    assert items[0]["candidate_name"] == "Test User"

    assert mock_repo.delete_resume("doc123") is True


def test_user_canonical_json_schema_structure():
    """Verify that the generated canonical JSON schema matches the user's required schema structure."""
    instance = ResumeCanonicalSchema(
        metadata=ResumeMetadataSchema(
            parser_version="1.0",
            status="pending",
            created_at=None,
        )
    )
    dumped = instance.model_dump()

    # Assert top-level sections
    expected_top_keys = {
        "profile",
        "education",
        "technical_skills",
        "projects",
        "experience",
        "certifications",
        "achievements",
        "metadata",
    }
    assert set(dumped.keys()) == expected_top_keys

    # Assert profile keys
    assert set(dumped["profile"].keys()) == {
        "full_name",
        "email",
        "phone",
        "location",
        "linkedin_url",
        "github_url",
        "portfolio_url",
        "summary",
    }

    # Assert education items structure
    edu_item = EducationItemSchema().model_dump()
    assert set(edu_item.keys()) == {
        "institution",
        "degree",
        "specialization",
        "start_year",
        "end_year",
        "cgpa",
        "cgpa_scale",
        "percentage",
        "class_10_percentage",
        "class_12_percentage",
        "backlogs",
    }

    # Assert technical_skills contains unified technologies array
    assert "technologies" in dumped["technical_skills"]
    assert isinstance(dumped["technical_skills"]["technologies"], list)

    # Assert projects item structure contains technologies
    proj_item = ProjectItemSchema().model_dump()
    assert set(proj_item.keys()) == {"title", "role", "description", "technologies"}

    # Assert experience item structure contains technologies
    exp_item = ExperienceItemSchema().model_dump()
    assert set(exp_item.keys()) == {
        "company",
        "role",
        "experience_type",
        "start_date",
        "end_date",
        "description",
        "technologies",
    }
    assert exp_item["experience_type"] == "unknown"

    # Assert certifications item structure
    cert_item = CertificationItemSchema().model_dump()
    assert set(cert_item.keys()) == {
        "name",
        "issuer",
        "issue_date",
        "expiry_date",
        "credential_url",
        "credential_id",
    }

    # Assert achievements item structure
    ach_item = AchievementItemSchema().model_dump()
    assert set(ach_item.keys()) == {"title", "description", "organization", "date"}

    # Assert metadata keys include required ones
    assert "parser_version" in dumped["metadata"]
    assert "status" in dumped["metadata"]
    assert "created_at" in dumped["metadata"]


# ==============================================================================
# SECTION 17 TESTS: Error Handling, AI Boundary, 401 Tests & Route Responses
# ==============================================================================
import openai
from starlette.testclient import TestClient
from app.main import app


def test_ai_401_unauthorized_raises_error_and_does_not_produce_empty_result():
    """Verify that OpenAI 401 Unauthorized raises AIExtractionError and does NOT fall back to empty lists or save."""
    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.find_by_file_hash.return_value = None
    service = ResumeService(repository=mock_repo)

    auth_error = openai.AuthenticationError(
        "Invalid API key provided",
        response=MagicMock(status_code=401),
        body={"error": {"code": "token_invalidated", "message": "Your API key has been invalidated."}}
    )

    with patch("openai.resources.chat.completions.Completions.create", side_effect=auth_error), \
         patch("app.core.config.settings.OPENAI_API_KEY", "invalid-key-sk-proj"), \
         patch.object(TextExtractor, "extract_text", return_value={
             "text": "Student Name\nstudent@example.com\nEDUCATION\nB.Tech CSE\nPROJECTS\nOnline E-Commerce System\nBuilt using React and Node.js.",
             "is_scanned": False,
             "page_count": 1,
             "file_hash": "hash_auth_fail_test",
             "file_size": 1500,
         }):
        with pytest.raises(AIExtractionError, match="401"):
            service.process_resume("student_resume.pdf", b"%PDF-mock")

    # Critical requirement: Database write must NOT be executed on AI failure
    assert not mock_repo.insert_resume.called


def test_ai_rate_limit_and_timeout_errors_raise_ai_extraction_error():
    """Verify that OpenAI RateLimitError and APITimeoutError raise AIExtractionError without silent fallback."""
    rate_error = openai.RateLimitError(
        "Rate limit exceeded",
        response=MagicMock(status_code=429),
        body={"error": {"message": "You exceeded your current quota."}}
    )
    with patch("openai.resources.chat.completions.Completions.create", side_effect=rate_error), \
         patch("app.core.config.settings.OPENAI_API_KEY", "mock-key"):
        with pytest.raises(AIExtractionError, match="429"):
            AIExtractor.extract_projects_and_experience(
                projects_text="Project Alpha in Python",
                experience_text="",
            )

    timeout_error = openai.APITimeoutError(request=MagicMock())
    with patch("openai.resources.chat.completions.Completions.create", side_effect=timeout_error), \
         patch("app.core.config.settings.OPENAI_API_KEY", "mock-key"):
        with pytest.raises(AIExtractionError, match="timed out"):
            AIExtractor.extract_projects_and_experience(
                projects_text="",
                experience_text="Software Intern at TechCorp",
            )


def test_ai_missing_key_when_ai_sections_present_raises_error():
    """Verify that if projects or experience text exists, missing OPENAI_API_KEY raises AIExtractionError."""
    with patch("app.core.config.settings.OPENAI_API_KEY", ""):
        with pytest.raises(AIExtractionError, match="OPENAI_API_KEY is not configured"):
            AIExtractor.extract_projects_and_experience(
                projects_text="IoT Weather Station using ESP32 and C++",
                experience_text="",
            )


def test_ai_legitimate_empty_sections_when_content_truly_absent():
    """Verify that when projects and experience are genuinely absent, AI is NOT called and empty lists succeed."""
    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.find_by_file_hash.return_value = None
    mock_repo.insert_resume.return_value = "doc_clean_no_projects"

    service = ResumeService(repository=mock_repo)

    # Resume with Profile, Education, Skills ONLY — NO Projects, NO Experience
    resume_text = """
    Priya Sharma
    priya.sharma@example.com | +91 9876543210
    Delhi, India

    EDUCATION
    B.Tech Computer Science and Engineering
    Delhi Technological University
    2020 - 2024 | CGPA: 9.0

    TECHNICAL SKILLS
    Languages: Python, Java, SQL
    """

    with patch("openai.resources.chat.completions.Completions.create") as mock_openai:
        with patch.object(TextExtractor, "extract_text", return_value={
            "text": resume_text,
            "is_scanned": False,
            "page_count": 1,
            "file_hash": "hash_no_projects_genuinely",
            "file_size": 1200,
        }):
            resp = service.process_resume("priya_clean.pdf", b"%PDF-mock")

        # OpenAI must NOT have been called
        assert not mock_openai.called
        assert resp.data.projects == []
        assert resp.data.experience == []
        assert resp.data.metadata.ai_called_for_projects_or_experience is False
        assert resp.success is True
        assert resp.saved is True


def test_composite_section_headers_and_field_cleaning():
    """Verify section detector and parser handle 'TECHNICAL & ACADEMIC PROJECTS' and

    'CERTIFICATIONS & TECHNICAL ACHIEVEMENTS' without losing data.
    """
    resume_text = """
    Alex Mercer
    Bangalore, India • +91 98765 01234 • alex.mercer@email.com
    linkedin.com/in/alexmercer-dev

    PROFESSIONAL SUMMARY
    Final-year B.Tech Information Technology undergraduate.

    EDUCATION
    Bachelor of Technology (B.Tech) in Information Technology
    National Institute of Technology • CGPA: 8.85 / 10.0
    Expected Graduation: 2027

    WORK EXPERIENCE
    Backend & Infrastructure Engineering Intern May 2026 - July 2026
    Nexus Distributed Labs Bangalore, India
    Implemented a cache-aside key-value invalidation pattern.

    TECHNICAL & ACADEMIC PROJECTS
    High-Throughput Distributed Event Engine
    Architected an append-only distributed log with continuous topic partition compaction.

    CERTIFICATIONS & TECHNICAL ACHIEVEMENTS
    Cloud & Infrastructure Accreditation: Validated foundational expertise in distributed cloud systems.
    Hackathon Winner: Awarded 1st place among 50 competing university teams for stream processing.
    """

    sections = SectionDetector.partition_resume(resume_text)
    assert len(sections["projects"]) > 0
    assert len(sections["experience"]) > 0
    assert len(sections["certifications"]) > 0

    # Test profile location does not capture preceding words or newlines
    profile = DeterministicParser.parse_profile(sections["header_profile"], sections["summary"], resume_text)
    assert profile.full_name == "Alex Mercer"
    assert profile.location == "Bangalore, India"

    # Test education separates institution from bullet and CGPA
    edu = DeterministicParser.parse_education(sections["education"])
    assert len(edu) == 1
    assert edu[0].institution == "National Institute of Technology"
    assert edu[0].cgpa == 8.85
    assert edu[0].cgpa_scale == 10.0

    # Test certs and achievements fallback separation
    certs = DeterministicParser.parse_certifications(sections["certifications"], fallback_text=sections["achievements"])
    achs = DeterministicParser.parse_achievements(sections["achievements"], fallback_text=sections["certifications"])
    assert any("Accreditation" in (c.name or "") for c in certs)
    assert any("Winner" in (a.title or "") for a in achs)


def test_api_route_handles_ai_extraction_error_with_502():
    """Verify POST /api/resumes/parse returns HTTP 502 Bad Gateway when AI extraction fails."""
    client = TestClient(app)

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Rahul Sharma\nEDUCATION\nB.Tech CSE\nPROJECTS\nMy Project in React and Python")
    pdf_bytes = doc.write()
    doc.close()

    auth_error = openai.AuthenticationError(
        "Invalid API key",
        response=MagicMock(status_code=401),
        body={"error": {"code": "token_invalidated", "message": "Invalid API key."}}
    )

    with patch("openai.resources.chat.completions.Completions.create", side_effect=auth_error), \
         patch("app.core.config.settings.OPENAI_API_KEY", "invalid-key"), \
         patch("app.core.config.settings.DEMO_MODE", True):
        files = {"file": ("test_resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
        resp = client.post("/api/resumes/parse", files=files)
        assert resp.status_code == 502
        assert "AI extraction failed" in resp.json()["detail"]


def test_api_route_handles_database_persistence_error_with_500():
    """Verify POST /api/resumes/parse returns HTTP 500 when database persistence fails."""
    client = TestClient(app)

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Rahul Sharma\nEDUCATION\nB.Tech CSE")
    pdf_bytes = doc.write()
    doc.close()

    with patch("app.repositories.resume_repository.ResumeRepository.insert_resume", side_effect=ResumeRepositoryError("Database connection lost")):
        files = {"file": ("test_resume.pdf", io.BytesIO(pdf_bytes), "application/pdf")}
        resp = client.post("/api/resumes/parse", files=files)
        assert resp.status_code == 500
        assert "Database persistence error" in resp.json()["detail"]


# ==============================================================================
# PHASE NINE REGRESSION & DYNAMIC TESTS (AC-01 to AC-12)
# ==============================================================================

def test_demo_pdf_four_projects_regression():
    """Verify demo.pdf regression: 4 projects preserved, exact skills preserved,

    no global skills contamination, empty experience, schema validity.
    """
    demo_resume_text = """
    Aayush Rathore
    aayush.rathore@example.com | +91 9876543210
    Indore, Madhya Pradesh, India
    linkedin.com/in/aayushrathore
    github.com/aayushrathore

    EDUCATION
    B.Tech in Computer Science and Engineering
    Medi-Caps University, Indore
    2022 - 2026 | CGPA: 8.5 / 10

    TECHNICAL SKILLS
    Languages: Python, JavaScript, SQL, HTML, CSS
    Frameworks: Django, Next.js
    Databases: MySQL, Supabase
    Tools: Git, GitHub, VS Code, Vercel
    Other Skills: Data Structures & Algorithms, Object-Oriented Programming, DBMS, Operating Systems, Computer Networks, NumPy, Pandas, Matplotlib, Streamlit

    PROJECTS
    ResQ-Now — Emergency Resource Connect Platform
    Tech Stack: Next.js, React, TypeScript, Supabase
    - Built a full-stack emergency resource platform connecting citizens, volunteers, NGOs, blood donors, and shelters.
    - Implemented authentication, location-based resource sharing, and real-time updates using Supabase.
    - Team of 3.

    Civic Saathi — AI-Powered Civic Issue Reporting Platform
    Tech Stack: Next.js, React, TypeScript, Gemini AI, Supabase, Google Maps API, Tailwind CSS
    - Built an AI-powered civic issue reporting platform with GPS-based location tracking and real-time complaint updates.
    - Integrated Gemini AI for complaint categorization, Google Maps visualization, and Supabase authentication with role-based dashboards.

    Sahyog — Artisan E-Commerce Platform
    Tech Stack: Django, Python
    - Developed an e-commerce platform connecting artisans with customers, with admin-managed artisan approval and product management.
    - Implemented artisan product listing, customer authentication, product browsing, shopping cart, and order management using Django ORM and CRUD operations.

    AI Fairness Auditor — Google Solution Challenge 2026
    Tech Stack: Python, Streamlit, scikit-learn, Pandas, NumPy
    - Developed an ML tool to detect demographic bias in datasets and machine learning models.
    - Built a Streamlit dashboard to visualize fairness metrics, including Disparate Impact Ratio.
    - Team of 3.

    ACHIEVEMENTS
    - Finalist in Google Solution Challenge 2026.
    - Solved 400+ problems on LeetCode.
    """

    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.find_by_file_hash.return_value = None
    mock_repo.insert_resume.return_value = "doc_demo_pdf_success_id"

    service = ResumeService(repository=mock_repo)

    mock_ai_projects = [
        {
            "title": "ResQ-Now — Emergency Resource Connect Platform",
            "role": None,
            "description": "Built a full-stack emergency resource platform connecting citizens, volunteers, NGOs, blood donors, and shelters. Implemented authentication, location-based resource sharing, and real-time updates using Supabase.",
            "technical_skills": ["Next.js", "React", "TypeScript", "Supabase"]
        },
        {
            "title": "Civic Saathi — AI-Powered Civic Issue Reporting Platform",
            "role": None,
            "description": "Built an AI-powered civic issue reporting platform with GPS-based location tracking and real-time complaint updates. Integrated Gemini AI for complaint categorization, Google Maps visualization, and Supabase authentication with role-based dashboards.",
            "technical_skills": ["Next.js", "React", "TypeScript", "Gemini AI", "Supabase", "Google Maps API", "Tailwind CSS"]
        },
        {
            "title": "Sahyog — Artisan E-Commerce Platform",
            "role": None,
            "description": "Developed an e-commerce platform connecting artisans with customers, with admin-managed artisan approval and product management. Implemented artisan product listing, customer authentication, product browsing, shopping cart, and order management using Django ORM and CRUD operations.",
            "technical_skills": ["Django", "Python"]
        },
        {
            "title": "AI Fairness Auditor — Google Solution Challenge 2026",
            "role": None,
            "description": "Developed an ML tool to detect demographic bias in datasets and machine learning models. Built a Streamlit dashboard to visualize fairness metrics, including Disparate Impact Ratio.",
            "technical_skills": ["Python", "Streamlit", "scikit-learn", "Pandas", "NumPy"]
        }
    ]

    mock_choice = MagicMock()
    mock_choice.message.content = json.dumps({
        "projects": mock_ai_projects,
        "experience": []
    })

    with patch("openai.resources.chat.completions.Completions.create", return_value=MagicMock(choices=[mock_choice])), \
         patch("app.core.config.settings.OPENAI_API_KEY", "mock-valid-key"), \
         patch.object(TextExtractor, "extract_text", return_value={
             "text": demo_resume_text,
             "is_scanned": False,
             "page_count": 1,
             "file_hash": "hash_demo_pdf_test",
             "file_size": 2500,
         }):
        res = service.process_resume("demo.pdf", b"%PDF-mock-bytes")

    assert res.success is True
    assert res.saved is True
    assert res.resume_id == "doc_demo_pdf_success_id"
    data = res.data

    # 1. Verify 4 distinct projects extracted
    assert len(data.projects) == 4
    project_titles = [p.title for p in data.projects]
    assert "ResQ-Now — Emergency Resource Connect Platform" in project_titles
    assert "Civic Saathi — AI-Powered Civic Issue Reporting Platform" in project_titles
    assert "Sahyog — Artisan E-Commerce Platform" in project_titles
    assert "AI Fairness Auditor — Google Solution Challenge 2026" in project_titles

    # 2. Verify project-specific technical skills and descriptions
    resq = next(p for p in data.projects if "ResQ-Now" in p.title)
    assert "Next.js" in resq.technical_skills
    assert "Supabase" in resq.technical_skills
    assert "shelters" in (resq.description or "")

    civic = next(p for p in data.projects if "Civic Saathi" in p.title)
    assert "Gemini AI" in civic.technical_skills
    assert "Google Maps API" in civic.technical_skills

    sahyog = next(p for p in data.projects if "Sahyog" in p.title)
    assert "Django" in sahyog.technical_skills
    assert "Python" in sahyog.technical_skills

    fairness = next(p for p in data.projects if "AI Fairness Auditor" in p.title)
    assert "Streamlit" in fairness.technical_skills
    assert "scikit-learn" in fairness.technical_skills

    # 3. Verify global technical skills matching dedicated skills section contains all extracted technologies
    expected_techs = [
        "Python", "JavaScript", "SQL", "HTML", "CSS", "Django", "Next.js",
        "MySQL", "Supabase", "Git", "GitHub", "VS Code", "Vercel",
        "Data Structures & Algorithms", "Object-Oriented Programming", "DBMS",
        "Operating Systems", "Computer Networks", "NumPy", "Pandas", "Matplotlib", "Streamlit"
    ]
    for tech in expected_techs:
        assert tech in data.technical_skills.technologies, f"Expected tech '{tech}' missing from technologies!"

    # 4. Critical check: Verify NO contamination of global skills from project descriptions or features
    contaminants = [
        "volunteers", "ngos", "blood donors", "shelters", "and shelters.",
        "based resource sharing", "gemini ai", "google maps api", "tailwind css",
        "artisan e-commerce platform", "implemented artisan product listing",
        "customer authentication", "shopping cart", "and crud operations.",
        "including disparate impact ratio."
    ]
    all_global_skills = [s.lower() for s in data.technical_skills.technologies]
    for c in contaminants:
        assert c not in all_global_skills, f"Contaminant '{c}' leaked into global technical skills!"

    # 5. Experience must be empty
    assert data.experience == []

    # 6. Schema validation & Database persistence
    assert mock_repo.insert_resume.called
    saved_doc = mock_repo.insert_resume.call_args[0][0]
    assert len(saved_doc["projects"]) == 4
    assert saved_doc["experience"] == []


def test_dynamic_resume_variations_ordering_and_missing_sections():
    """Verify dynamic support for reordered sections, single project, and missing experience."""
    reordered_text = """
    Priya Sharma
    priya.sharma@example.com
    Bengaluru, India

    PROJECTS
    Automated Video Transcoding Pipeline
    Technologies: Go, FFmpeg, AWS S3, Docker
    Designed distributed video transcoding workers scaling to 100 concurrent jobs.

    TECHNICAL SKILLS
    Languages: Go, Python
    Tools: Docker, Kubernetes, Git

    EDUCATION
    B.E. in Information Science
    BMS College of Engineering, Bangalore
    2020 - 2024
    """

    mock_repo = MagicMock(spec=ResumeRepository)
    mock_repo.find_by_file_hash.return_value = None
    mock_repo.insert_resume.return_value = "doc_reordered_id"
    service = ResumeService(repository=mock_repo)

    mock_choice = MagicMock()
    mock_choice.message.content = json.dumps({
        "projects": [
            {
                "title": "Automated Video Transcoding Pipeline",
                "role": "Backend Engineer",
                "description": "Designed distributed video transcoding workers scaling to 100 concurrent jobs.",
                "technical_skills": ["Go", "FFmpeg", "AWS S3", "Docker"]
            }
        ],
        "experience": []
    })

    with patch("openai.resources.chat.completions.Completions.create", return_value=MagicMock(choices=[mock_choice])), \
         patch("app.core.config.settings.OPENAI_API_KEY", "mock-valid-key"), \
         patch.object(TextExtractor, "extract_text", return_value={
             "text": reordered_text,
             "is_scanned": False,
             "page_count": 1,
             "file_hash": "hash_reordered_test",
             "file_size": 1800,
         }):
        res = service.process_resume("priya_resume.pdf", b"%PDF-mock-bytes")

    assert res.success is True
    assert res.data.profile.full_name == "Priya Sharma"
    assert res.data.profile.location == "Bengaluru, India"
    assert len(res.data.projects) == 1
    assert res.data.projects[0].title == "Automated Video Transcoding Pipeline"
    assert res.data.projects[0].technical_skills == ["Go", "FFmpeg", "AWS S3", "Docker"]
    assert res.data.experience == []
    assert res.data.technical_skills.programming_languages == ["Go", "Python"]
    assert res.data.technical_skills.tools == ["Docker", "Kubernetes", "Git"]
    assert len(res.data.education) == 1
    assert res.data.education[0].institution == "BMS College of Engineering, Bangalore"



